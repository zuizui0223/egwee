from __future__ import annotations

import io
import urllib.request
from docx import Document

URL = "https://static-content.springer.com/esm/art%3A10.1186%2F1472-6785-13-10/MediaObjects/12898_2012_239_MOESM1_ESM.docx"
UA = "Mozilla/5.0 egwee-magnolia-ml011-schema/1.0"


def main() -> None:
    req = urllib.request.Request(URL, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        payload = r.read()
        print(f"MAGNOLIA_SUPP_ACCESS status={r.status} type={r.headers.get('content-type')!r} bytes={len(payload)}")
    doc = Document(io.BytesIO(payload))
    print(f"MAGNOLIA_SUPP_SCHEMA paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
    for i, p in enumerate(doc.paragraphs, 1):
        text = " ".join(p.text.split())
        if text and not any(ch.isdigit() for ch in text):
            print(f"MAGNOLIA_SUPP_LABEL paragraph={i} text={text!r}")
    for ti, table in enumerate(doc.tables, 1):
        rows, cols = len(table.rows), len(table.columns)
        header = []
        if rows:
            header = [" ".join(c.text.split()) for c in table.rows[0].cells]
        print(f"MAGNOLIA_SUPP_TABLE table={ti} rows={rows} cols={cols} header={header!r}")
    print("MAGNOLIA_ML011_SUPP_SCHEMA_GATE PASS no_numeric_data_rows_printed")


if __name__ == "__main__":
    main()
