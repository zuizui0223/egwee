from __future__ import annotations
import csv, io, urllib.request
from pathlib import Path
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
REG=ROOT/'evidence/meta_extraction/multilayer_cluster_registry_v1.csv'
RES=ROOT/'manuscript/MAGNOLIA_ML011_CLUSTER_RECOVERY_RESULT.md'
URL='https://static-content.springer.com/esm/art%3A10.1186%2F1472-6785-13-10/MediaObjects/12898_2012_239_MOESM1_ESM.docx'
EXPECTED={'Y':85,'T':97,'A':46,'B':16,'C':23,'F':4}
def main():
    with REG.open(newline='',encoding='utf-8') as f: rows=list(csv.DictReader(f))
    by={r['cluster_id']:r for r in rows}
    assert by['ML009']['species']=='Pistacia lentiscus'
    assert by['ML010']['species']=='Penstemon hirsutus'
    m=by['ML011']
    assert m['study_ids']=='PS002' and m['species']=='Magnolia stellata'
    assert m['fragmentation_contrast']=='source_defined_summed_adult_genet_basal_area'
    assert int(m['n_admissible_primary_effects'])==0
    assert m['covariance_status']=='blocked_before_effect_calculation'
    assert m['cluster_status']=='source_population_size_values_not_recoverable'
    adm=[r for r in rows if r['cluster_status']=='admissible_multilayer_cluster']
    assert {r['cluster_id'] for r in adm}=={'ML001','ML002','ML003','ML014'}
    assert sum(int(r['n_admissible_primary_effects']) for r in adm)==11
    req=urllib.request.Request(URL,headers={'User-Agent':'Mozilla/5.0 egwee-magnolia-ml011-audit/1.0'})
    with urllib.request.urlopen(req,timeout=60) as r: payload=r.read()
    doc=Document(io.BytesIO(payload)); assert len(doc.tables)==1
    t=doc.tables[0]; assert len(t.rows)==8 and len(t.columns)==21
    got={}
    for row in t.rows[2:]:
        vals=[' '.join(c.text.split()) for c in row.cells]
        got[vals[0]]=int(vals[1])
    assert got==EXPECTED, (got,EXPECTED)
    text=RES.read_text(encoding='utf-8')
    assert 'No C/F effect or covariance was calculated' in text
    assert '85, 97, 46, 16, 23, 4' in text
    print('Magnolia ML011 closure: PASS; source exposure mismatch audited; ML011 adds 0 effects; current primary denominator=4/11')
if __name__=='__main__': main()
