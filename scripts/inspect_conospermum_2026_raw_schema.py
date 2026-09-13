from __future__ import annotations

import io
import json
import urllib.parse
import urllib.request
import zipfile

from openpyxl import load_workbook

DOI = "10.5061/dryad.95x69p907"
ROOT = "https://datadryad.org/api/v2"
UA = "egwee-conospermum-2026-raw-schema/1.2"
TARGETS = {"Seedlings_scoring.xlsx", "Paternity_dataset.xlsx", "README.md"}


def dataset_archive() -> bytes:
    key = urllib.parse.quote(f"doi:{DOI}", safe="")
    url = f"{ROOT}/datasets/{key}/download"
    req = urllib.request.Request(
        url,
        headers={"User-Agent": UA, "Accept": "application/zip,application/octet-stream,*/*"},
    )
    with urllib.request.urlopen(req, timeout=120) as r:
        payload = r.read()
    if not payload.startswith(b"PK"):
        raise RuntimeError(f"dataset download did not return ZIP bytes: prefix={payload[:16]!r}")
    return payload


def archive_member_map(zf: zipfile.ZipFile) -> dict[str, str]:
    mapping = {}
    for name in zf.namelist():
        leaf = name.rstrip("/").split("/")[-1]
        if leaf in TARGETS:
            mapping[leaf] = name
    return mapping


def main() -> None:
    payload = dataset_archive()
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        members = archive_member_map(zf)
        missing = TARGETS - set(members)
        if missing:
            raise RuntimeError(f"missing archive targets: {sorted(missing)}; members={zf.namelist()!r}")

        readme = zf.read(members["README.md"]).decode("utf-8", errors="replace")
        print("CONOSPERMUM_2026_README " + json.dumps(readme, ensure_ascii=False))

        for name in ("Seedlings_scoring.xlsx", "Paternity_dataset.xlsx"):
            workbook_bytes = zf.read(members[name])
            wb = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
            print(f"CONOSPERMUM_2026_WORKBOOK name={name!r} sheets={wb.sheetnames!r} bytes={len(workbook_bytes)}")
            for ws in wb.worksheets:
                rows = ws.iter_rows(values_only=True)
                first = next(rows, tuple())
                print(
                    "CONOSPERMUM_2026_SHEET "
                    f"workbook={name!r} sheet={ws.title!r} rows={ws.max_row} cols={ws.max_column} "
                    f"header={list(first)!r}"
                )
    print("CONOSPERMUM_2026_RAW_SCHEMA_GATE PASS values_beyond_headers_not_printed")


if __name__ == "__main__":
    main()
