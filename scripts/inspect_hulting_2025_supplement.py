from __future__ import annotations

import io
import urllib.parse
import urllib.request

from docx import Document

DOI = "10.1111/1365-2745.14452"
FILE = "jec14452-sup-0001-Supinfo.docx"
URL = "https://onlinelibrary.wiley.com/action/downloadSupplement?doi=" + urllib.parse.quote(DOI, safe="") + "&file=" + FILE
UA = "Mozilla/5.0 egwee-hulting-2025-supplement/1.0"


def clean(x: str) -> str:
    return " ".join((x or "").split())


def main() -> None:
    req = urllib.request.Request(URL, headers={"User-Agent": UA, "Accept": "application/vnd.openxmlformats-officedocument.wordprocessingml.document,*/*"})
    with urllib.request.urlopen(req, timeout=60) as response:
        payload = response.read()
    if not payload.startswith(b"PK"):
        prefix = payload[:120].decode("utf-8", errors="replace")
        raise RuntimeError(f"supplement did not return docx bytes: {prefix!r}")

    doc = Document(io.BytesIO(payload))
    print(f"HULTING_2025_SUPP bytes={len(payload)} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} url={URL!r}")
    for i, p in enumerate(doc.paragraphs, 1):
        text = clean(p.text)
        if text and ("Table S" in text or "Figure S" in text):
            print(f"HULTING_2025_CAPTION {i:03d}: {text}")

    for ti, table in enumerate(doc.tables, 1):
        print(f"HULTING_2025_TABLE_BEGIN {ti} rows={len(table.rows)} cols={max((len(r.cells) for r in table.rows), default=0)}")
        for ri, row in enumerate(table.rows, 1):
            vals = [clean(c.text) for c in row.cells]
            print(f"HULTING_2025_TABLE {ti} ROW {ri}: " + " | ".join(vals))
        print(f"HULTING_2025_TABLE_END {ti}")


if __name__ == "__main__":
    main()
