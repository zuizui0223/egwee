from __future__ import annotations

import csv
import io
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook

# File IDs and DOI/version identities are locked by the source-qualification probe.
PUBLIC_DOIS = {
    "paternity_2026": "10.5061/dryad.95x69p907",
    "reproduction_2019": "10.5061/dryad.4cg374r",
}
PUBLIC_FILES = {
    "paternity_2026": {
        "Paternity_dataset_unformatted.xlsx": 4702034,
        "Seedlings_scoring_unformatted.xlsx": 4702035,
    },
    "reproduction_2019": {
        "fruit_and_seed_set.csv": 155093,
    },
}


def payload_looks_real(name: str, payload: bytes) -> bool:
    if not payload:
        return False
    if name.lower().endswith(".xlsx"):
        return payload[:2] == b"PK"
    if name.lower().endswith(".csv"):
        prefix = payload[:256].lower()
        return b"<html" not in prefix and b"<!doctype" not in prefix and b"," in prefix
    return True


def download_bytes(label: str, name: str, file_id: int) -> bytes:
    doi = PUBLIC_DOIS[label]
    landing = f"https://datadryad.org/dataset/{quote(f'doi:{doi}', safe='')}"
    stream = f"https://datadryad.org/downloads/file_stream/{file_id}"
    with tempfile.TemporaryDirectory() as tmp:
        jar = str(Path(tmp) / "cookies.txt")
        # Establish the same landing-page session used by a browser before
        # following the exact href printed in the public dataset HTML.
        landing_proc = subprocess.run(
            [
                "curl", "-L", "--fail", "--silent", "--show-error",
                "-A", "Mozilla/5.0 egwee-conospermum-2026-audit/1.0",
                "-c", jar,
                landing,
            ],
            check=False,
            capture_output=True,
        )
        if landing_proc.returncode != 0:
            raise RuntimeError(
                f"Dryad landing session failed label={label} rc={landing_proc.returncode} "
                f"stderr={landing_proc.stderr.decode('utf-8', errors='replace')}"
            )
        proc = subprocess.run(
            [
                "curl", "-L", "--fail", "--silent", "--show-error",
                "--retry", "2", "--retry-delay", "1",
                "-A", "Mozilla/5.0 egwee-conospermum-2026-audit/1.0",
                "-e", landing,
                "-b", jar,
                "-c", jar,
                stream,
            ],
            check=False,
            capture_output=True,
        )
        if proc.returncode != 0:
            raise RuntimeError(
                f"Dryad browser-session download failed label={label} file_id={file_id} "
                f"rc={proc.returncode} stderr={proc.stderr.decode('utf-8', errors='replace')}"
            )
        payload = proc.stdout
    print(
        f"CONO2026_SESSION_DOWNLOAD label={label!r} file_id={file_id} "
        f"bytes={len(payload)} magic={payload[:24].hex()!r}"
    )
    if not payload_looks_real(name, payload):
        prefix = payload[:500].decode("utf-8", errors="replace")
        raise RuntimeError(
            f"Dryad public href returned non-file content for {name} (id={file_id}); "
            f"prefix={prefix!r}"
        )
    return payload


def inspect_xlsx(label: str, name: str, payload: bytes) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / name
        path.write_bytes(payload)
        wb = load_workbook(path, read_only=True, data_only=True)
        print(f"CONO2026_XLSX label={label!r} name={name!r} sheets={wb.sheetnames!r} bytes={len(payload)}")
        for ws in wb.worksheets:
            print(f"CONO2026_SHEET file={name!r} name={ws.title!r} rows={ws.max_row} cols={ws.max_column}")
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i > 12:
                    break
                values = ["" if v is None else str(v) for v in row[:36]]
                print(f"CONO2026_PREVIEW file={name!r} sheet={ws.title!r} row={i} values={values!r}")


def inspect_csv(label: str, name: str, payload: bytes) -> None:
    text = payload.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    fields = reader.fieldnames or []
    print(f"CONO2026_CSV label={label!r} name={name!r} rows={len(rows)} fields={fields!r}")
    for row in rows[:8]:
        print(f"CONO2026_CSV_PREVIEW name={name!r} row={row!r}")

    exposure_fields = [
        field for field in fields
        if any(token in field.lower() for token in ("connect", "isolat", "area", "size", "population", "display"))
    ]
    print(f"CONO2026_EXPOSURE_FIELDS name={name!r} fields={exposure_fields!r}")
    for field in exposure_fields:
        values = sorted({str(r.get(field, "")).strip() for r in rows if str(r.get(field, "")).strip()})
        print(f"CONO2026_FIELD_VALUES field={field!r} n_unique={len(values)} sample={values[:40]!r}")

    population_fields = [f for f in fields if "pop" in f.lower()]
    connectivity_fields = [f for f in fields if any(t in f.lower() for t in ("connect", "isolat"))]
    print(
        f"CONO2026_EXPOSURE_SCHEMA population_fields={population_fields!r} "
        f"connectivity_fields={connectivity_fields!r}"
    )
    assert population_fields, fields


def main() -> None:
    for label, files in PUBLIC_FILES.items():
        for name, file_id in files.items():
            payload = download_bytes(label, name, file_id)
            if name.lower().endswith(".xlsx"):
                inspect_xlsx(label, name, payload)
            elif name.lower().endswith(".csv"):
                inspect_csv(label, name, payload)
            else:
                raise AssertionError(name)

    print("CONOSPERMUM_2026_SOURCE_INSPECTION PASS")


if __name__ == "__main__":
    main()
