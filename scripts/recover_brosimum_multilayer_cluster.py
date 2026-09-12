from __future__ import annotations

import csv
import json
import math
import statistics as stats
import tempfile
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ARTICLE_ID = 22130177
TARGET = "Datos-Brosimum.xlsx"
API = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}"
SITE_SNAPSHOT = ROOT / "evidence/meta_extraction/PS004_brosimum_site_table_v1.csv"
VALID_HABITATS = {"CON", "FRA"}

# Published site-level multilocus correlated-paternity values (Table 3), already
# independently audited in evidence/meta_extraction/PS004_brosimum_extraction_v1.csv.
# The raw workbook abbreviates Cuixmala as CUI.
RP_BY_POP = {
    "CHA": 0.106,
    "CAR": 0.164,
    "CUI": 0.107,
    "ASE": 0.246,
    "TEC": 0.208,
    "ZAP": 0.191,
}


def fetch_json(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": "egwee-public-data-audit/1.0"})
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.load(response)


def download(url: str, path: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "egwee-public-data-audit/1.0"})
    with urllib.request.urlopen(request, timeout=60) as response, path.open("wb") as fh:
        fh.write(response.read())


def number(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {".", "NA", "N/A", "nan"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def hedges_g(fragmented: list[float], reference: list[float]) -> tuple[float, float]:
    """Hedges g with metafor escalc(SMD, vtype='LS') sampling variance."""
    n1, n2 = len(fragmented), len(reference)
    assert n1 >= 2 and n2 >= 2
    m1, m2 = stats.mean(fragmented), stats.mean(reference)
    s1, s2 = stats.stdev(fragmented), stats.stdev(reference)
    df = n1 + n2 - 2
    pooled = math.sqrt(((n1 - 1) * s1**2 + (n2 - 1) * s2**2) / df)
    d = (m1 - m2) / pooled
    j = 1 - 3 / (4 * df - 1)
    g = j * d
    variance = 1 / n1 + 1 / n2 + g**2 / (2 * (n1 + n2))
    return g, variance


def pearson(x: list[float], y: list[float]) -> float:
    assert len(x) == len(y) and len(x) >= 3
    mx, my = stats.mean(x), stats.mean(y)
    dx = [v - mx for v in x]
    dy = [v - my for v in y]
    return sum(a * b for a, b in zip(dx, dy)) / math.sqrt(
        sum(a * a for a in dx) * sum(b * b for b in dy)
    )


def group_centered(values: dict[str, float], habitat_by_pop: dict[str, str], pops: list[str]) -> list[float]:
    means = {
        habitat: stats.mean(values[p] for p in pops if habitat_by_pop[p] == habitat)
        for habitat in VALID_HABITATS
    }
    return [values[p] - means[habitat_by_pop[p]] for p in pops]


def load_workbook_from_figshare() -> object:
    payload = fetch_json(API)
    matches = [f for f in payload.get("files", []) if f.get("name") == TARGET]
    assert len(matches) == 1, matches
    item = matches[0]
    assert int(item["id"]) == 39338195
    assert item.get("computed_md5") == "049db31f8ebcec42c4c44ebe4a6af70a"
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / TARGET
        download(item["download_url"], path)
        assert path.stat().st_size == int(item["size"])
        wb = load_workbook(path, read_only=False, data_only=True)
    return wb


def recover_vigor(wb: object) -> tuple[dict[str, str], dict[str, dict[str, float]], dict[str, int]]:
    ws = wb["Vigor"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    required = {"Habitat", "Pop", "MadID", "NumHojasTot", "AreFoliarTot", "Heigh", "TPDW"}
    assert required <= set(idx), header

    endpoints = ["NumHojasTot", "AreFoliarTot", "Heigh", "TPDW"]
    maternal: dict[tuple[str, str, str], dict[str, list[float]]] = defaultdict(
        lambda: defaultdict(list)
    )
    habitat_by_pop: dict[str, str] = {}

    for row in rows[1:]:
        habitat = str(row[idx["Habitat"]]).strip() if row[idx["Habitat"]] is not None else ""
        pop = str(row[idx["Pop"]]).strip() if row[idx["Pop"]] is not None else ""
        mother = str(row[idx["MadID"]]).strip() if row[idx["MadID"]] is not None else ""
        if habitat not in VALID_HABITATS or not pop or not mother:
            continue
        if pop in habitat_by_pop:
            assert habitat_by_pop[pop] == habitat
        habitat_by_pop[pop] = habitat
        for endpoint in endpoints:
            value = number(row[idx[endpoint]])
            if value is not None:
                maternal[(habitat, pop, mother)][endpoint].append(value)

    pop_maternal_values: dict[str, dict[str, list[float]]] = defaultdict(lambda: defaultdict(list))
    mothers_per_pop: dict[str, set[str]] = defaultdict(set)
    for (habitat, pop, mother), endpoint_values in maternal.items():
        mothers_per_pop[pop].add(mother)
        for endpoint, values in endpoint_values.items():
            if values:
                pop_maternal_values[pop][endpoint].append(stats.mean(values))

    site_means: dict[str, dict[str, float]] = defaultdict(dict)
    for pop, endpoint_values in pop_maternal_values.items():
        for endpoint, values in endpoint_values.items():
            site_means[endpoint][pop] = stats.mean(values)

    mother_counts = {pop: len(ids) for pop, ids in mothers_per_pop.items()}
    return habitat_by_pop, dict(site_means), mother_counts


def recover_genetic_ho(wb: object) -> tuple[dict[str, str], dict[str, dict[str, float]], dict[str, dict[str, int]]]:
    ws = wb["Genetic__data"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    assert header[:4] == ["Habitat", "Pop", "Develomental_stage", "ID"], header[:4]
    locus_cols = list(range(4, len(header), 2))
    assert len(locus_cols) == 8, locus_cols

    heterozygous: dict[tuple[str, str], list[int]] = defaultdict(list)
    n_individuals: dict[tuple[str, str], set[str]] = defaultdict(set)
    habitat_by_pop: dict[str, str] = {}

    for row in rows[1:]:
        habitat = str(row[0]).strip() if row[0] is not None else ""
        pop = str(row[1]).strip() if row[1] is not None else ""
        stage = str(row[2]).strip() if row[2] is not None else ""
        individual = str(row[3]).strip() if row[3] is not None else ""
        if habitat not in VALID_HABITATS or not pop or not stage or not individual:
            continue
        if pop in habitat_by_pop:
            assert habitat_by_pop[pop] == habitat
        habitat_by_pop[pop] = habitat
        n_individuals[(pop, stage)].add(individual)
        for col in locus_cols:
            a1 = number(row[col])
            a2 = number(row[col + 1])
            if a1 is None or a2 is None:
                continue
            heterozygous[(pop, stage)].append(int(a1 != a2))

    stages = sorted({stage for (_, stage) in heterozygous})
    by_stage: dict[str, dict[str, float]] = {stage: {} for stage in stages}
    counts: dict[str, dict[str, int]] = {stage: {} for stage in stages}
    for (pop, stage), indicators in heterozygous.items():
        assert indicators
        by_stage[stage][pop] = stats.mean(indicators)
        counts[stage][pop] = len(n_individuals[(pop, stage)])
    return habitat_by_pop, by_stage, counts


def check_snapshot(
    habitat: dict[str, str],
    vigor: dict[str, dict[str, float]],
    genetic: dict[str, dict[str, float]],
    mother_counts: dict[str, int],
    genetic_counts: dict[str, dict[str, int]],
) -> None:
    with SITE_SNAPSHOT.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert len(rows) == 6
    by_pop = {r["population"]: r for r in rows}
    assert set(by_pop) == set(habitat) == set(RP_BY_POP)
    for pop, row in by_pop.items():
        assert row["habitat"] == habitat[pop]
        assert abs(float(row["C_paternity_rp"]) - RP_BY_POP[pop]) < 1e-12
        assert abs(float(row["F_TPDW_site_mean"]) - vigor["TPDW"][pop]) < 1e-9
        assert int(row["n_maternal_trees_vigor"]) == mother_counts[pop]
        assert abs(float(row["G_adult_Ho_raw_diagnostic"]) - genetic["AD"][pop]) < 1e-9
        assert abs(float(row["G_progeny_Ho_raw_diagnostic"]) - genetic["PR"][pop]) < 1e-9
        assert int(row["n_adult_genotype_rows"]) == genetic_counts["AD"][pop]
        assert int(row["n_progeny_genotype_rows"]) == genetic_counts["PR"][pop]


def main() -> None:
    wb = load_workbook_from_figshare()
    vigor_habitat, vigor, mother_counts = recover_vigor(wb)
    genetic_habitat, genetic, genetic_counts = recover_genetic_ho(wb)

    assert vigor_habitat == genetic_habitat
    habitat = vigor_habitat
    pops = sorted(habitat)
    assert set(pops) == set(RP_BY_POP), (pops, sorted(RP_BY_POP))
    assert set(genetic) == {"AD", "PR"}
    check_snapshot(habitat, vigor, genetic, mother_counts, genetic_counts)

    print(f"BROSIMUM_RECOVERY pops={pops!r}")
    print(f"BROSIMUM_RECOVERY habitat={habitat!r}")
    print(f"BROSIMUM_RECOVERY maternal_tree_counts={mother_counts!r}")
    print(f"BROSIMUM_RECOVERY genetic_individual_counts={genetic_counts!r}")

    fragmented = [p for p in pops if habitat[p] == "FRA"]
    reference = [p for p in pops if habitat[p] == "CON"]
    assert len(fragmented) == len(reference) == 3

    f_values = vigor["TPDW"]
    f_g, f_var = hedges_g([f_values[p] for p in fragmented], [f_values[p] for p in reference])
    print(f"BROSIMUM_EFFECT layer=F endpoint=TPDW g={f_g:.12f} var={f_var:.12f}")

    # Genetic values are reconstructed for QA only. They are not promoted until
    # the raw-file summaries reconcile with the publication's Table 2 summaries.
    for stage, values in sorted(genetic.items()):
        g, variance = hedges_g([values[p] for p in fragmented], [values[p] for p in reference])
        print(f"BROSIMUM_DIAGNOSTIC layer=G stage={stage!r} g={g:.12f} var={variance:.12f}")

    c_support = {p: -RP_BY_POP[p] for p in pops}
    c_g_raw, c_var = hedges_g([RP_BY_POP[p] for p in fragmented], [RP_BY_POP[p] for p in reference])
    c_g = -c_g_raw
    print(f"BROSIMUM_EFFECT layer=C endpoint=rp oriented_g={c_g:.12f} var={c_var:.12f}")

    vectors = {"C_rp_support": c_support, "F_TPDW": f_values}
    names = list(vectors)
    centered = {name: group_centered(vectors[name], habitat, pops) for name in names}
    for i, a in enumerate(names):
        for b in names[i:]:
            r = pearson(centered[a], centered[b])
            print(f"BROSIMUM_CORR a={a!r} b={b!r} r={r:.12f}")

    print("BROSIMUM_RECOVERY snapshot=PASS; C/F site-level cluster recoverable; G remains QA-pending")


if __name__ == "__main__":
    main()
