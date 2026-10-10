from __future__ import annotations

import csv
import io
import json
import urllib.request
import zipfile
from pathlib import PurePosixPath

from openpyxl import load_workbook
from bs4 import BeautifulSoup

COLLECTION_ID = 3300914
API_ROOT = "https://api.figshare.com/v2"
UA = "egwee-leptonychia-public-data-audit/1.0"
MAX_SCHEMA_BYTES = 25_000_000


def get_bytes(url: str) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()


def get_json(url: str):
    return json.loads(get_bytes(url).decode("utf-8-sig"))


def csv_schema(raw: bytes) -> dict:
    text = raw.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return {
        "rows_including_header": len(rows),
        "columns": rows[0] if rows else [],
    }


def xlsx_schema(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    out = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out.append(
            {
                "sheet": ws.title,
                "rows": ws.max_row,
                "columns": ["" if x is None else str(x).strip() for x in first],
            }
        )
    return {"sheets": out}


def html_schema(raw: bytes) -> dict:
    soup = BeautifulSoup(raw, "html.parser")
    headings = []
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
        text = " ".join(tag.stripped_strings)
        if text:
            headings.append(text)

    tables = []
    for i, table in enumerate(soup.find_all("table"), start=1):
        rows = table.find_all("tr")
        first = rows[0] if rows else None
        columns = []
        if first is not None:
            columns = [" ".join(cell.stripped_strings) for cell in first.find_all(["th", "td"])]
        caption = ""
        cap = table.find("caption")
        if cap is not None:
            caption = " ".join(cap.stripped_strings)
        row_values = []
        for row in rows:
            row_values.append(
                [" ".join(cell.stripped_strings) for cell in row.find_all(["th", "td"])]
            )

        context = ""
        prev = table.find_previous(["p", "div"])
        if prev is not None:
            context = " ".join(prev.stripped_strings)
            if len(context) > 1200:
                context = context[-1200:]

        tables.append(
            {
                "index": i,
                "caption": caption,
                "rows": len(rows),
                "columns": columns,
                "context": context,
                "row_values": row_values,
            }
        )
    return {"headings": headings, "tables": tables}


def main() -> None:
    articles = get_json(
        f"{API_ROOT}/collections/{COLLECTION_ID}/articles?page=1&page_size=1000"
    )
    assert isinstance(articles, list) and articles, "Figshare collection returned no articles"

    print(
        "LEPTONYCHIA_FIGSHARE_COLLECTION "
        f"collection_id={COLLECTION_ID} n_articles={len(articles)}"
    )

    for item in articles:
        aid = item.get("id")
        assert aid is not None
        meta = get_json(f"{API_ROOT}/articles/{aid}")
        files = meta.get("files", [])
        print(
            "LEPTONYCHIA_FIGSHARE_ARTICLE "
            f"article_id={aid} title={meta.get('title')!r} n_files={len(files)}"
        )

        for f in files:
            name = f.get("name") or ""
            size = int(f.get("size") or 0)
            url = f.get("download_url")
            print(
                "LEPTONYCHIA_FIGSHARE_FILE "
                f"article_id={aid} id={f.get('id')} name={name!r} "
                f"size={size} md5={f.get('computed_md5')!r}"
            )
            if not url or size <= 0 or size > MAX_SCHEMA_BYTES:
                continue

            suffix = PurePosixPath(name).suffix.casefold()
            if suffix not in {".csv", ".tsv", ".xlsx", ".zip", ".txt", ".htm", ".html"}:
                continue

            raw = get_bytes(url)
            assert raw, (aid, name)

            if suffix in {".csv", ".tsv", ".txt"}:
                schema = csv_schema(raw)
                print(
                    "LEPTONYCHIA_SCHEMA "
                    f"name={name!r} type=text rows={schema['rows_including_header']} "
                    f"columns={schema['columns']!r}"
                )
            elif suffix == ".xlsx":
                schema = xlsx_schema(raw)
                print(
                    "LEPTONYCHIA_SCHEMA "
                    f"name={name!r} type=xlsx sheets={schema['sheets']!r}"
                )
            elif suffix == ".zip":
                with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                    names = [n for n in zf.namelist() if not n.endswith("/")]
                print(
                    "LEPTONYCHIA_SCHEMA "
                    f"name={name!r} type=zip members={names!r}"
                )
            elif suffix in {".htm", ".html"}:
                schema = html_schema(raw)
                print(
                    "LEPTONYCHIA_SCHEMA "
                    f"name={name!r} type=html headings={schema['headings']!r}"
                )
                for table in schema["tables"]:
                    print(
                        "LEPTONYCHIA_HTML_TABLE "
                        f"index={table['index']} caption={table['caption']!r} "
                        f"context={table['context']!r} rows={table['row_values']!r}"
                    )

    print("LEPTONYCHIA_FIGSHARE_SCHEMA_OK outcomes_summarized=false")


if __name__ == "__main__":
    main()
