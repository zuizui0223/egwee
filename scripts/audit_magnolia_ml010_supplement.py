from __future__ import annotations
import io, urllib.request
from docx import Document
URL='https://static-content.springer.com/esm/art%3A10.1186%2F1472-6785-13-10/MediaObjects/12898_2012_239_MOESM1_ESM.docx'
UA='Mozilla/5.0 egwee-magnolia-ml010-audit/1.0'
EXPECTED={'Y':85,'T':97,'A':46,'B':16,'C':23,'F':4}
def main():
    req=urllib.request.Request(URL,headers={'User-Agent':UA})
    with urllib.request.urlopen(req,timeout=60) as r: payload=r.read()
    doc=Document(io.BytesIO(payload)); assert len(doc.tables)==1
    t=doc.tables[0]; assert len(t.rows)==8 and len(t.columns)==21
    got={}
    for row in t.rows[2:]:
        vals=[' '.join(c.text.split()) for c in row.cells]
        pop=vals[0]; got[pop]=int(vals[1])
    assert got==EXPECTED, (got,EXPECTED)
    # Source Methods defines the model exposure as summed adult-genet basal area;
    # these small integer supplement values are the census-count representation,
    # so they are audited but not admitted as the locked exposure.
    print(f'Magnolia ML010 supplement audit: PASS; supplement population-size column={got}; no C/F effect calculated from this proxy')
if __name__=='__main__': main()
