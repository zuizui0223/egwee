from __future__ import annotations

import hashlib
import io
import json
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evidence/meta_extraction/phase2_cf01_hass_supplement_schema_v1.json"
STATUS = ROOT / "manuscript/PHASE2_CF01_HASS_SUPPLEMENT_SCHEMA_2026-09-24.md"

COLLECTION_ID = 3994269
TARGET = "rspb20172242supp3.xlsx"
UA = "Mozilla/5.0 egwee-hass-2018-schema/1.0"


def get_bytes(url: str, accept: str = "*/*") -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": accept})
    with urllib.request.urlopen(req, timeout=90) as r:
        return r.read()


def get_json(url: str):
    return json.loads(get_bytes(url, "application/json").decode("utf-8-sig"))


def figshare_target() -> tuple[str, bytes] | None:
    try:
        items = get_json(f"https://api.figshare.com/v2/collections/{COLLECTION_ID}/articles")
    except Exception:
        return None
    if not isinstance(items, list):
        return None
    for item in items:
        aid = item.get("id")
        if aid is None:
            continue
        try:
            article = get_json(f"https://api.figshare.com/v2/articles/{aid}")
        except Exception:
            continue
        for f in article.get("files", []):
            name = (f.get("name") or "").strip()
            url = f.get("download_url")
            if name == TARGET and url:
                try:
                    return url, get_bytes(url)
                except Exception:
                    pass
    return None


def download_target() -> tuple[str | None, bytes | None, list[str]]:
    errors: list[str] = []
    hit = figshare_target()
    if hit is not None:
        return hit[0], hit[1], errors

    urls = [
        f"https://pmc.ncbi.nlm.nih.gov/articles/PMC5829195/bin/{TARGET}",
        (
            "https://royalsocietypublishing.org/action/downloadSupplement"
            "?doi=10.1098%2Frspb.2017.2242&file=" + TARGET
        ),
    ]
    for url in urls:
        try:
            raw = get_bytes(url)
            head = raw[:512].lstrip().lower()
            if not raw or b"<html" in head or b"<!doctype html" in head:
                errors.append(f"{url}: non-workbook response bytes={len(raw)}")
                continue
            return url, raw, errors
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
    return None, None, errors


def workbook_schema(raw: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    out = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = ["" if x is None else str(x).strip() for x in first]
        low = [c.casefold() for c in columns]
        out.append({
            "sheet": ws.title,
            "rows": ws.max_row,
            "columns": columns,
            "landscape_id_hints": [
                columns[i] for i, x in enumerate(low)
                if "landscape" in x or x in {"land", "landid", "land_id"}
            ],
            "field_id_hints": [
                columns[i] for i, x in enumerate(low)
                if "field" in x or x in {"site", "siteid", "site_id"}
            ],
            "exposure_hints": [
                columns[i] for i, x in enumerate(low)
                if ("border" in x and ("dens" in x or "length" in x))
                or "config" in x
            ],
            "wild_bee_hints": [
                columns[i] for i, x in enumerate(low)
                if ("wild" in x and "bee" in x) or "wildbee" in x
            ],
            "seed_set_hints": [
                columns[i] for i, x in enumerate(low)
                if "seed" in x or "pod" in x
            ],
        })
    return out


def main() -> None:
    url, raw, errors = download_target()
    access = raw is not None

    result = {
        "schema_version": 1,
        "candidate": "CFTQ0084",
        "programme_id": "P2_CF01_HASS_WEUROPE_2018",
        "article_doi": "10.1098/rspb.2017.2242",
        "supplement_collection_doi": "10.6084/m9.figshare.c.3994269",
        "target_file": TARGET,
        "access_status": "public_s3_recoverable" if access else "public_s3_access_blocked",
        "download_url_used": url,
        "download_errors": errors,
        "file_sha256": hashlib.sha256(raw).hexdigest() if raw else None,
        "file_bytes": len(raw) if raw else None,
        "sheets": workbook_schema(raw) if raw else [],
        "effect_outcomes_opened": False,
        "numeric_effects_calculated": False,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    STATUS.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    lines = [
        "# CFTQ0084 Hass western-Europe S3 schema gate — 2026-09-24",
        "",
        "- article: doi:10.1098/rspb.2017.2242",
        "- supplement collection: doi:10.6084/m9.figshare.c.3994269",
        f"- target file: `{TARGET}`",
        f"- public S3 access: **{'recoverable' if access else 'blocked'}**",
        "- numerical EGWEE effect calculation opened: **false**",
        "",
    ]
    if access:
        lines += [
            "## Schema",
            "",
            "Only workbook sheet names, dimensions and header fields were inspected in this gate.",
            "No response values, correlations, coefficients, p-values or effect directions were calculated.",
            "",
        ]
        for s in result["sheets"]:
            lines += [
                f"### {s['sheet']}",
                "",
                f"- rows: {s['rows']}",
                f"- columns: {', '.join(s['columns'])}",
                f"- landscape-ID hints: {', '.join(s['landscape_id_hints']) or 'none'}",
                f"- field-ID hints: {', '.join(s['field_id_hints']) or 'none'}",
                f"- field-border/configuration hints: {', '.join(s['exposure_hints']) or 'none'}",
                f"- wild-bee hints: {', '.join(s['wild_bee_hints']) or 'none'}",
                f"- seed/pod hints: {', '.join(s['seed_set_hints']) or 'none'}",
                "",
            ]
        lines += [
            "## Next gate",
            "",
            "Proceed only if S3 exposes a reproducible field-to-landscape mapping, landscape-level",
            "field-border density, wild-bee abundance and radish seeds-per-pod response on a common",
            "landscape frame. The frozen recovery contract requires landscape, not field, as n.",
            "",
        ]
    else:
        lines += [
            "## Access STOP",
            "",
            "The article declares S3/S4 public, but the automated public routes tested here did not",
            "yield the S3 workbook. This is an access boundary, not an ecological result. Do not",
            "replace S3 with digitised figures or published significance summaries.",
            "",
        ]

    STATUS.write_text("\n".join(lines), encoding="utf-8")
    print(
        "PHASE2_CF01_HASS_SCHEMA_"
        + ("OK " if access else "STOP ")
        + f"access={str(access).lower()} sheets={len(result['sheets'])} effects_opened=false"
    )


if __name__ == "__main__":
    main()
