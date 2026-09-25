from __future__ import annotations

import csv
import io
import urllib.request
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evidence/meta_extraction/phase2_cf01_sevenello_supp_table_s1_v1.csv"
STATUS = ROOT / "manuscript/PHASE2_CF01_SEVENELLO_SUPP_TABLE_S1_2026-09-25.md"

URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1007%2Fs10980-026-02305-2/MediaObjects/"
    "10980_2026_2305_MOESM1_ESM.docx"
)
UA = "egwee-sevenello-supp-s1/1.0"


def download() -> bytes:
    req = urllib.request.Request(URL, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=90) as response:
        return response.read()


def clean(text: str) -> str:
    return " ".join((text or "").replace("\n", " ").replace("\r", " ").split())


def main() -> None:
    raw = download()
    doc = Document(io.BytesIO(raw))
    assert doc.tables, "supplement contains no tables"

    # Outcome-blind: inspect only the first supplementary table (Table S1),
    # which the manuscript identifies as the study-site / focal-species layout.
    table = doc.tables[0]
    rows = [[clean(cell.text) for cell in row.cells] for row in table.rows]
    assert len(rows) >= 2

    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]

    # Preserve the source table literally enough to audit aliases/occupancy.
    header = [x if x else f"column_{i+1}" for i, x in enumerate(padded[0])]
    # DOCX merged cells can repeat header names; make them unique mechanically.
    seen: dict[str, int] = {}
    unique_header = []
    for h in header:
        seen[h] = seen.get(h, 0) + 1
        unique_header.append(h if seen[h] == 1 else f"{h}__{seen[h]}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(unique_header)
        writer.writerows(padded[1:])

    lines = [
        "# CFTQ0001 Sevenello supplementary Table S1 audit — 2026-09-25",
        "",
        "Only the **first supplementary table (Table S1)** was opened. Later supplementary",
        "tables containing model/results information were not inspected in this gate.",
        "",
        f"- DOCX bytes: **{len(raw)}**",
        f"- Table S1 rows including header: **{len(rows)}**",
        f"- Table S1 columns: **{width}**",
        "",
        "## Table S1 transcription",
        "",
    ]
    for row in padded:
        lines.append("- " + " | ".join(row))

    lines += [
        "",
        "## Gate",
        "",
        "Use Table S1 only to reconcile reserve/remnant aliases, crop-side sampling units,",
        "edge/core transects and focal-species occupancy. Do not use later supplementary",
        "result tables before the candidate-specific recovery contract is frozen.",
        "",
    ]
    STATUS.parent.mkdir(parents=True, exist_ok=True)
    STATUS.write_text("\n".join(lines), encoding="utf-8")

    print(
        "PHASE2_CF01_SEVENELLO_SUPP_S1_OK "
        f"rows={len(rows)} cols={width} later_tables_opened=false"
    )


if __name__ == "__main__":
    main()
