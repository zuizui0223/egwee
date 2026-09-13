from __future__ import annotations

import io
import urllib.request
from docx import Document

URL = "https://ars.els-cdn.com/content/image/1-s2.0-S000632072030882X-mmc1.docx"


def main() -> None:
    req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0 egwee-conospermum-audit/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        payload = r.read()
    doc = Document(io.BytesIO(payload))
    print(f"CONOSPERMUM_SUPP_SUMMARY bytes={len(payload)} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")
    for i, p in enumerate(doc.paragraphs, start=1):
        text = " ".join(p.text.split())
        if text:
            print(f"CONOSPERMUM_SUPP_PARA {i:03d} {text}")
    for ti, table in enumerate(doc.tables, start=1):
        print(f"CONOSPERMUM_SUPP_TABLE_BEGIN {ti}")
        for ri, row in enumerate(table.rows, start=1):
            vals = [" ".join(cell.text.split()) for cell in row.cells]
            print(f"CONOSPERMUM_SUPP_TABLE {ti} ROW {ri}: " + " | ".join(vals))
        print(f"CONOSPERMUM_SUPP_TABLE_END {ti}")


if __name__ == "__main__":
    main()
