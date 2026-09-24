from __future__ import annotations

import hashlib
import io
import json
import re
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SUMMARY = ROOT / "evidence/meta_extraction/phase2_cf01_bdffp_dryad_schema_v1.json"
STATUS = ROOT / "manuscript/PHASE2_CF01_BDFFP_DRYAD_SCHEMA_2026-09-24.md"

DOI = "10.5061/dryad.612jm640h"
DATASET_API = "https://datadryad.org/api/v2/datasets/" + urllib.parse.quote("doi:" + DOI, safe="")
PUBLIC_ROOT = "https://datadryad.org"
UA = "Mozilla/5.0 egwee-bdffp-schema/1.1"

EXPECTED = {
    "density.rich.data.xlsx",
    "dispersed.matrix.xlsx",
    "functional.diversity.xlsx",
    "undispersed.matrix.xlsx",
}

STRUCTURAL_TERMS = (
    "plot", "site", "fragment", "size", "ranch", "forest", "treatment",
    "habitat", "class", "area", "control",
)
RESPONSE_TERMS = (
    "dispers", "undispers", "density", "seed", "rich", "divers", "abund",
)


def request_bytes(url: str) -> tuple[int, str, bytes]:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=90) as resp:
        return resp.status, resp.headers.get("content-type", ""), resp.read()


def get_json(url: str) -> dict[str, Any]:
    status, ctype, payload = request_bytes(url)
    if status != 200:
        raise RuntimeError(f"GET {url} -> {status}")
    try:
        return json.loads(payload.decode("utf-8-sig"))
    except Exception as exc:
        raise RuntimeError(f"expected JSON from {url}; type={ctype!r} bytes={len(payload)}") from exc


def absolute(href: str) -> str:
    return urllib.parse.urljoin(PUBLIC_ROOT, href)


def iter_links(obj: Any):
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key == "href" and isinstance(value, str):
                yield value
            yield from iter_links(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from iter_links(value)


def embedded_items(obj: dict[str, Any], token: str) -> list[dict[str, Any]]:
    emb = obj.get("_embedded", {})
    if isinstance(emb, dict):
        for key, value in emb.items():
            if token.casefold() in key.casefold() and isinstance(value, list):
                return [x for x in value if isinstance(x, dict)]
    return []


def resolve_version(dataset: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    hrefs = list(iter_links(dataset))
    exact = [h for h in hrefs if re.search(r"/api/v2/versions/\d+/?$", h)]
    if exact:
        url = absolute(exact[0])
        return url, get_json(url)

    collections = [h for h in hrefs if "/api/v2/datasets/" in h and "/versions" in h]
    if not collections:
        raise RuntimeError("Dryad dataset metadata has no version relation")
    page = get_json(absolute(collections[0]))
    versions = embedded_items(page, "version")
    if not versions:
        raise RuntimeError("Dryad versions collection is empty")
    versions.sort(key=lambda x: int(x.get("versionNumber", x.get("id", 0)) or 0))
    chosen = versions[-1]
    links = list(iter_links(chosen))
    exact = [h for h in links if re.search(r"/api/v2/versions/\d+/?$", h)]
    if exact:
        url = absolute(exact[0])
        return url, get_json(url)
    vid = chosen.get("id")
    if vid is None:
        raise RuntimeError("cannot resolve Dryad version id")
    url = f"{PUBLIC_ROOT}/api/v2/versions/{vid}"
    return url, get_json(url)


def resolve_files(version_url: str, version: dict[str, Any]) -> list[dict[str, Any]]:
    hrefs = list(iter_links(version))
    candidates = [h for h in hrefs if "/files" in h]
    listing = get_json(absolute(candidates[0])) if candidates else get_json(version_url.rstrip("/") + "/files")
    files = embedded_items(listing, "file")
    if files:
        return files
    for value in listing.values():
        if isinstance(value, list) and value and all(isinstance(x, dict) for x in value):
            return value
    raise RuntimeError(f"could not identify files in Dryad response keys={list(listing)}")


def file_name(item: dict[str, Any]) -> str:
    for key in ("path", "name", "filename"):
        value = item.get(key)
        if isinstance(value, str) and value:
            return value.rsplit("/", 1)[-1]
    return ""


def candidate_download_urls(item: dict[str, Any]) -> list[str]:
    urls: list[str] = []
    for href in iter_links(item):
        low = href.casefold()
        if "download" in low or "file_stream" in low:
            urls.append(absolute(href))
    fid = item.get("id")
    if fid is not None:
        urls.extend([
            f"{PUBLIC_ROOT}/stash/downloads/file_stream/{fid}",
            f"{PUBLIC_ROOT}/downloads/file_stream/{fid}",
            f"{PUBLIC_ROOT}/api/v2/files/{fid}/download",
        ])
    return list(dict.fromkeys(urls))


def download_file(item: dict[str, Any]) -> tuple[str, bytes]:
    errors: list[str] = []
    for url in candidate_download_urls(item):
        try:
            status, ctype, payload = request_bytes(url)
            if status == 200 and payload:
                head = payload[:512].lstrip().casefold()
                if b"<html" in head or b"<!doctype html" in head:
                    errors.append(f"{url}: html response")
                    continue
                return url, payload
            errors.append(f"{url}: status={status} type={ctype!r} bytes={len(payload)}")
        except Exception as exc:
            errors.append(f"{url}: {type(exc).__name__}: {exc}")
    raise RuntimeError("; ".join(errors))


def workbook_schema(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = ["" if x is None else str(x).strip() for x in first]
        structural = [
            c for c in columns if c and any(t in c.casefold() for t in STRUCTURAL_TERMS)
        ]
        response = [
            c for c in columns if c and any(t in c.casefold() for t in RESPONSE_TERMS)
        ]

        structural_values: dict[str, list[str]] = {}
        col_index = {c: i for i, c in enumerate(columns)}
        if structural:
            found = {c: set() for c in structural}
            for row in ws.iter_rows(min_row=2, values_only=True):
                for c in structural:
                    i = col_index[c]
                    if i < len(row) and row[i] is not None:
                        found[c].add(str(row[i]).strip())
            structural_values = {c: sorted(vals)[:100] for c, vals in found.items()}

        sheets.append({
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "columns": columns,
            "structural_columns": structural,
            "response_columns": response,
            "structural_values": structural_values,
        })
    return {"sheets": sheets}


def main() -> None:
    dataset = get_json(DATASET_API)
    version_url, version = resolve_version(dataset)
    files = resolve_files(version_url, version)
    by_name = {file_name(item): item for item in files}
    missing = sorted(EXPECTED - set(by_name))
    if missing:
        raise RuntimeError(f"missing expected Dryad files: {missing}; available={sorted(by_name)}")

    inspected = {}
    download_urls = {}
    for short in sorted(EXPECTED):
        url, raw = download_file(by_name[short])
        download_urls[short] = url
        inspected[short] = {
            "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "schema": workbook_schema(raw),
        }

    schema_text = json.dumps(inspected, ensure_ascii=False).casefold()
    plot_hint = "plot" in schema_text
    fragment_hint = "fragment" in schema_text or "size" in schema_text
    dispersed_hint = "dispers" in schema_text
    undispersed_hint = "undispers" in schema_text

    result = {
        "schema_version": 2,
        "candidate": "CFTQ0065",
        "programme_id": "P2_CF01_BDFFP_SEED_RAIN_2020",
        "dataset_doi": DOI,
        "article_doi": "10.1002/eap.2093",
        "dryad_dataset_id": dataset.get("id"),
        "dryad_version_id": version.get("id"),
        "dryad_version_number": version.get("versionNumber"),
        "files": inspected,
        "download_urls_used": download_urls,
        "schema_hints": {
            "plot_identifier_present": plot_hint,
            "fragment_exposure_present": fragment_hint,
            "dispersed_seed_endpoint_present": dispersed_hint,
            "undispersed_seed_endpoint_present": undispersed_hint,
        },
        "effect_outcomes_opened": False,
        "numeric_effects_calculated": False,
        "next_gate": (
            "verify the eleven source plots and common dispersed/undispersed density frame, "
            "then execute the frozen direct C-F recovery contract"
        ),
    }

    SUMMARY.parent.mkdir(parents=True, exist_ok=True)
    STATUS.parent.mkdir(parents=True, exist_ok=True)
    SUMMARY.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    lines = [
        "# CFTQ0065 BDFFP Dryad schema gate — 2026-09-24",
        "",
        "## Source",
        "",
        "- article: doi:10.1002/eap.2093",
        f"- public data: doi:{DOI}",
        f"- Dryad version: **{version.get('versionNumber')}** (id={version.get('id')})",
        "- recovery contract: `CF01_BDFFP_SEED_RAIN_2020_RECOVERY_CONTRACT.md`",
        "",
        "This gate uses Dryad's public metadata plus per-file public download routes. It inspects",
        "workbook structure and exposure/unit identifiers only; it does not summarize response values.",
        "",
        "## Archive result",
        "",
        f"- public files indexed: **{len(files)}**",
        "- all four declared workbooks present: **yes**",
        f"- plot identifier hint present: **{'yes' if plot_hint else 'no'}**",
        f"- fragment exposure hint present: **{'yes' if fragment_hint else 'no'}**",
        f"- dispersed-seed endpoint hint present: **{'yes' if dispersed_hint else 'no'}**",
        f"- undispersed-seed endpoint hint present: **{'yes' if undispersed_hint else 'no'}**",
        "- numerical EGWEE effect calculation opened: **false**",
        "",
        "## Workbook schemas",
        "",
    ]
    for short in sorted(inspected):
        lines.append(f"### {short}")
        lines.append("")
        for sheet in inspected[short]["schema"]["sheets"]:
            lines.append(
                f"- sheet `{sheet['sheet']}`: rows={sheet['max_row']}; "
                f"columns={', '.join(sheet['columns'])}"
            )
            if sheet["structural_columns"]:
                lines.append("- structural columns: " + ", ".join(sheet["structural_columns"]))
                for col, vals in sheet["structural_values"].items():
                    lines.append(f"- `{col}` structural values: {', '.join(vals)}")
        lines.append("")

    lines += [
        "## Gate",
        "",
        "Proceed to numerical recovery only if the public workbooks support all of:",
        "",
        "1. exactly the source plot frame can be reconstructed without using traps as n;",
        "2. fragment/control status is joined before opening endpoint values;",
        "3. dispersed and undispersed density are available on the same plot frame;",
        "4. the frozen all-fragment-versus-continuous contrast can be calculated without",
        "   selecting fragment sizes or response subsets from outcomes.",
        "",
        "Otherwise retain the programme as design-valid but quantitatively blocked.",
        "",
    ]
    STATUS.write_text("\n".join(lines), encoding="utf-8")

    print(
        "PHASE2_CF01_BDFFP_DRYAD_SCHEMA_OK "
        f"indexed={len(files)} plot_hint={str(plot_hint).lower()} "
        f"fragment_hint={str(fragment_hint).lower()} "
        f"dispersed_hint={str(dispersed_hint).lower()} "
        f"undispersed_hint={str(undispersed_hint).lower()} "
        "effects_opened=false"
    )


if __name__ == "__main__":
    main()
