from __future__ import annotations

import csv
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SUMMARY = ROOT / "evidence/meta_extraction/phase2_cf01_zurich_envidat_schema_v1.json"
STATUS = ROOT / "manuscript/PHASE2_CF01_ZURICH_ENVIDAT_SCHEMA_2026-09-24.md"

EXPECTED_SUFFIXES = {
    "data_description.xlsx",
    "protocol_english.pdf",
    "protocol_german.pdf",
    "garden_site_coordinates.csv",
    "taxa_checklist.csv",
    "raw_sampling_data.xlsx",
    "individual_traits.csv",
    "species_temporal_flower_visitation_matrix.csv",
    "daucus_carota_seed_set.csv",
    "onobrychis_viciifolia_fruit_set.csv",
    "raphanus_sativus_fruit_set.csv",
    "raphanus_sativus_seed_set.csv",
    "symphytum_officinale_fruit_set.csv",
    "symphytum_officinale_seed_set.csv",
    "figure_1.R",
    "figure_2.R",
    "figure_3.R",
    "figure_4.R",
    "table_2.R",
    "README.txt",
}


def basename(name: str) -> str:
    return name.rstrip("/").split("/")[-1]


def csv_schema(zf: zipfile.ZipFile, name: str) -> dict:
    raw = zf.read(name)
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        header = []
    n_rows = sum(1 for _ in reader)
    return {
        "path": name,
        "bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "columns": header,
        "data_rows": n_rows,
    }


def xlsx_schema(zf: zipfile.ZipFile, name: str) -> dict:
    raw = zf.read(name)
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        columns = [str(x).strip() if x is not None else "" for x in first]
        sheets.append({
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "columns": columns,
        })
    return {
        "path": name,
        "bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "sheets": sheets,
    }


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("usage: inspect_cf01_zurich_envidat_schema.py DATASET.zip")
    path = Path(sys.argv[1])
    assert path.is_file(), path

    with zipfile.ZipFile(path) as zf:
        names = [n for n in zf.namelist() if not n.endswith("/")]
        base = {basename(n): n for n in names}
        missing = sorted(EXPECTED_SUFFIXES - set(base))
        assert not missing, missing

        csv_files = {}
        for short in sorted(x for x in EXPECTED_SUFFIXES if x.endswith(".csv")):
            csv_files[short] = csv_schema(zf, base[short])

        xlsx_files = {}
        for short in ("data_description.xlsx", "raw_sampling_data.xlsx"):
            xlsx_files[short] = xlsx_schema(zf, base[short])

        readme = zf.read(base["README.txt"]).decode("utf-8-sig", errors="replace")

    exposure_tokens = ("impervious", "500", "garden", "urban")
    response_tokens = ("visit", "pollin", "seed", "fruit")
    schema_text = json.dumps({"csv": csv_files, "xlsx": xlsx_files}, ensure_ascii=False).lower()
    exposure_schema_hint = all(tok in (schema_text + readme.lower()) for tok in ("garden",))
    has_impervious_hint = "impervious" in (schema_text + readme.lower())
    response_schema_hint = any(tok in schema_text for tok in response_tokens)

    summary = {
        "schema_version": 1,
        "candidate": "CFTQ0018",
        "dataset_doi": "10.16904/envidat.676",
        "article_doi": "10.1016/j.dib.2025.112013",
        "archive_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "archive_bytes": path.stat().st_size,
        "file_count": len(names),
        "expected_files_present": True,
        "csv_files": csv_files,
        "xlsx_files": xlsx_files,
        "schema_hints": {
            "garden_identifier_or_context_present": exposure_schema_hint,
            "impervious_surface_term_present": has_impervious_hint,
            "pollination_or_reproductive_response_terms_present": response_schema_hint,
        },
        "numeric_outcome_summaries_calculated": False,
        "effect_outcomes_opened": False,
        "next_gate": (
            "identify exact garden key and 500m impervious-surface field, then verify that visitation "
            "and species seed/fruit records join at garden level before opening response values"
        ),
    }
    SUMMARY.write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    raw_sheets = xlsx_files["raw_sampling_data.xlsx"]["sheets"]
    STATUS.write_text(
        f"""# CFTQ0018 Zurich EnviDat schema gate — 2026-09-24

## Source

Open EnviDat dataset: doi:10.16904/envidat.676, linked to Reji Chacko et al. (2025),
doi:10.1016/j.dib.2025.112013.

The archive is inspected at schema level only before numerical response extraction.

## Archive result

- archive bytes: **{path.stat().st_size}**
- files found: **{len(names)}**
- all expected repository files present: **yes**
- raw field workbook sheets: **{len(raw_sheets)}**
- explicit impervious-surface term found in schema/README: **{"yes" if has_impervious_hint else "no"}**
- pollination/reproductive response terms found in schema: **{"yes" if response_schema_hint else "no"}**

The repository contains separate site, raw field, visitation/trait, and six species-level
seed/fruit-set files. No response means, effect directions, correlations, p-values, or
fragmentation effects were calculated in this gate.

## Admission question

The next gate is structural:

1. identify the exact garden identifier shared across exposure, visitation and reproductive files;
2. identify the source 500-m impervious-surface field;
3. verify that each candidate species has a common garden frame for I and F;
4. keep garden as the independent landscape unit;
5. only then freeze one I endpoint and one F endpoint before calculating a gradient effect.

If impervious exposure cannot be joined at garden level, CFTQ0018 remains descriptive rather than
being rescued with another urbanisation variable.
""",
        encoding="utf-8",
    )

    print(
        "PHASE2_CF01_ZURICH_SCHEMA_OK "
        f"files={len(names)} raw_sheets={len(raw_sheets)} "
        f"impervious_hint={has_impervious_hint} response_hint={response_schema_hint} "
        "outcomes_opened=false"
    )


if __name__ == "__main__":
    main()
