from __future__ import annotations

import json
import math
import statistics as stats
import tempfile
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ARTICLE_ID = 22130177
TARGET = "Datos-Brosimum.xlsx"
API = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}"
VALID_HABITATS = {"CON", "FRA"}

# Publication Table 2 habitat-level H_O, recorded before this reconciliation.
TABLE2_HO = {
    ("CON", "AD"): 0.62,
    ("FRA", "AD"): 0.63,
    ("CON", "PR"): 0.58,
    ("FRA", "PR"): 0.59,
}


def number(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {".", "NA", "N/A", "nan"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def fetch_json(url: str) -> dict:
    req = urllib.request.Request(url, headers={"User-Agent": "egwee-brosimum-genetic-qa/1.0"})
    with urllib.request.urlopen(req, timeout=30) as response:
        return json.load(response)


def download(url: str, path: Path) -> None:
    req = urllib.request.Request(url, headers={"User-Agent": "egwee-brosimum-genetic-qa/1.0"})
    with urllib.request.urlopen(req, timeout=60) as response, path.open("wb") as fh:
        fh.write(response.read())


def load_workbook_public() -> object:
    payload = fetch_json(API)
    matches = [f for f in payload.get("files", []) if f.get("name") == TARGET]
    assert len(matches) == 1, matches
    item = matches[0]
    assert int(item["id"]) == 39338195
    assert item.get("computed_md5") == "049db31f8ebcec42c4c44ebe4a6af70a"
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / TARGET
        download(item["download_url"], path)
        assert path.stat().st_size == int(item["size"])
        wb = load_workbook(path, read_only=False, data_only=True)
    return wb


def extract_calls(wb: object):
    ws = wb["Genetic__data"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    assert header[:4] == ["Habitat", "Pop", "Develomental_stage", "ID"], header[:4]
    locus_cols = list(range(4, len(header), 2))
    assert len(locus_cols) == 8

    calls: dict[tuple[str, str, str, int], list[int]] = defaultdict(list)
    individuals: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    habitats: dict[str, str] = {}

    for row in rows[1:]:
        habitat = str(row[0]).strip() if row[0] is not None else ""
        pop = str(row[1]).strip() if row[1] is not None else ""
        stage = str(row[2]).strip() if row[2] is not None else ""
        individual = str(row[3]).strip() if row[3] is not None else ""
        if habitat not in VALID_HABITATS or not pop or stage not in {"AD", "PR"} or not individual:
            continue
        if pop in habitats:
            assert habitats[pop] == habitat
        habitats[pop] = habitat
        individuals[(habitat, pop, stage)].add(individual)

        for locus_i, col in enumerate(locus_cols):
            a1 = number(row[col])
            a2 = number(row[col + 1])
            if a1 is None or a2 is None:
                continue
            calls[(habitat, pop, stage, locus_i)].append(int(a1 != a2))

    return calls, individuals, habitats


def mean_locus_ho(groups: list[list[int]]) -> float:
    locus_hos = [stats.mean(g) for g in groups if g]
    assert locus_hos
    return stats.mean(locus_hos)


def pooled_call_ho(groups: list[list[int]]) -> float:
    flat = [x for g in groups for x in g]
    assert flat
    return stats.mean(flat)


def main() -> None:
    wb = load_workbook_public()
    calls, individuals, habitats = extract_calls(wb)
    pops = sorted(habitats)
    assert len(pops) == 6
    assert sum(habitats[p] == "CON" for p in pops) == 3
    assert sum(habitats[p] == "FRA" for p in pops) == 3

    site_locus_mean: dict[tuple[str, str], float] = {}
    site_pooled: dict[tuple[str, str], float] = {}

    for pop in pops:
        habitat = habitats[pop]
        for stage in ("AD", "PR"):
            groups = [calls[(habitat, pop, stage, locus)] for locus in range(8)]
            assert all(groups), (pop, stage, [len(g) for g in groups])
            lm = mean_locus_ho(groups)
            pc = pooled_call_ho(groups)
            site_locus_mean[(pop, stage)] = lm
            site_pooled[(pop, stage)] = pc
            counts = [len(g) for g in groups]
            print(
                "BROSIMUM_SITE_HO "
                f"pop={pop} habitat={habitat} stage={stage} "
                f"locus_mean={lm:.12f} pooled_calls={pc:.12f} "
                f"locus_n={counts!r} n_individuals={len(individuals[(habitat,pop,stage)])}"
            )

    # Habitat-level reconstruction A: pool all individuals inside habitat at each locus,
    # compute H_O per locus, then average loci equally (ARLEQUIN-like estimator).
    habitat_locus_mean: dict[tuple[str, str], float] = {}
    habitat_pooled: dict[tuple[str, str], float] = {}
    mean_of_sites: dict[tuple[str, str], float] = {}

    for habitat in ("CON", "FRA"):
        hpops = [p for p in pops if habitats[p] == habitat]
        for stage in ("AD", "PR"):
            locus_groups = []
            for locus in range(8):
                merged = []
                for pop in hpops:
                    merged.extend(calls[(habitat, pop, stage, locus)])
                locus_groups.append(merged)
            lm = mean_locus_ho(locus_groups)
            pc = pooled_call_ho(locus_groups)
            ms = stats.mean(site_locus_mean[(p, stage)] for p in hpops)
            habitat_locus_mean[(habitat, stage)] = lm
            habitat_pooled[(habitat, stage)] = pc
            mean_of_sites[(habitat, stage)] = ms
            target = TABLE2_HO[(habitat, stage)]
            print(
                "BROSIMUM_HABITAT_HO "
                f"habitat={habitat} stage={stage} table2={target:.2f} "
                f"locus_mean_pooled_individuals={lm:.12f} "
                f"pooled_calls={pc:.12f} mean_site_locus_mean={ms:.12f} "
                f"delta_locus_mean={lm-target:+.12f} "
                f"delta_mean_sites={ms-target:+.12f}"
            )

    # The publication reports two decimal places. A valid reconstruction must round
    # to the four Table-2 H_O values under one consistent estimator.
    locus_round_match = all(
        round(habitat_locus_mean[key] + 1e-12, 2) == target
        for key, target in TABLE2_HO.items()
    )
    pooled_round_match = all(
        round(habitat_pooled[key] + 1e-12, 2) == target
        for key, target in TABLE2_HO.items()
    )
    mean_sites_round_match = all(
        round(mean_of_sites[key] + 1e-12, 2) == target
        for key, target in TABLE2_HO.items()
    )

    print(
        "BROSIMUM_TABLE2_MATCH "
        f"locus_mean_pooled_individuals={str(locus_round_match).lower()} "
        f"pooled_calls={str(pooled_round_match).lower()} "
        f"mean_site_locus_mean={str(mean_sites_round_match).lower()}"
    )

    # Fail closed if no single transparent reconstruction matches all four targets.
    assert locus_round_match or mean_sites_round_match, (
        "Neither locus-weighted habitat reconstruction nor mean site-level locus H_O "
        "reproduces all four Table-2 H_O values at reported precision."
    )


if __name__ == "__main__":
    main()
