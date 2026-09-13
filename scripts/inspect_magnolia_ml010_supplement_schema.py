from __future__ import annotations
import io, urllib.request
from docx import Document
URL='https://static-content.springer.com/esm/art%3A10.1186%2F1472-6785-13-10/MediaObjects/12898_2012_239_MOESM1_ESM.docx'
UA='Mozilla/5.0 egwee-magnolia-ml010-schema/1.0'
def main():
    req=urllib.request.Request(URL,headers={'User-Agent':UA})
    with urllib.request.urlopen(req,timeout=60) as r:
        payload=r.read(); print(f'MAGNOLIA_ML010_SUPP_ACCESS status={r.status} bytes={len(payload)}')
    doc=Document(io.BytesIO(payload)); print(f'MAGNOLIA_ML010_SCHEMA paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}')
    assert len(doc.tables)==1
    t=doc.tables[0]; header=[' '.join(c.text.split()) for c in t.rows[0].cells]
    print(f'MAGNOLIA_ML010_TABLE rows={len(t.rows)} cols={len(t.columns)} header={header!r}')
    assert any('Population size' in x for x in header)
    assert any('Source population of pollen' in x for x in header)
    print('MAGNOLIA_ML010_SCHEMA_GATE PASS')
if __name__=='__main__': main()
