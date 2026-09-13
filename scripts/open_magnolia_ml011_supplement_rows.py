from __future__ import annotations

import io
import urllib.request
from docx import Document

URL = "https://static-content.springer.com/esm/art%3A10.1186%2F1472-6785-13-10/MediaObjects/12898_2012_239_MOESM1_ESM.docx"
UA = "Mozilla/5.0 egwee-magnolia-ml011-open/1.0"


def main() -> None:
    req = urllib.request.Request(URL, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=60) as r:
        payload = r.read()
    doc = Document(io.BytesIO(payload))
    if len(doc.tables) != 1:
        raise RuntimeError(f"expected one table, found {len(doc.tables)}")
    table = doc.tables[0]
    for ri, row in enumerate(table.rows, 1):
        vals = [" ".join(cell.text.split()) for cell in row.cells]
        print(f"MAGNOLIA_ML011_ROW {ri:02d} {vals!r}")


if __name__ == "__main__":
    main()
