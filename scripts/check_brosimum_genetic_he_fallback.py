from __future__ import annotations

import json
import math
import statistics as stats
import tempfile
import urllib.request
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ARTICLE_ID = 22130177
TARGET = "Datos-Brosimum.xlsx"
API = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}"
RULE = ROOT / "manuscript/BROSIMUM_GPAIR_RECONCILIATION_RULE_2026-09-19.md"
STATUS = ROOT / "manuscript/PHASE2_ML002_BROSIMUM_GPAIR_RECONCILIATION_2026-09-19.md"
VALID_HABITATS = {"CON", "FRA"}
STAGES = ("AD", "PR")

TABLE2_HE = {
    ("CON", "AD"): 0.65,
    ("FRA", "AD"): 0.63,
    ("CON", "PR"): 0.59,
    ("FRA", "PR"): 0.60,
}


def number(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    if not text or text in {".", "NA", "N/A", "nan"}:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def fetch_json(url: str) -> dict:
    req = urllib.request.Request(url, headers={"User-Agent": "egwee-brosimum-he-audit/1.0"})
    with urllib.request.urlopen(req, timeout=30) as response:
        return json.load(response)


def download(url: str, path: Path) -> None:
    req = urllib.request.Request(url, headers={"User-Agent": "egwee-brosimum-he-audit/1.0"})
    with urllib.request.urlopen(req, timeout=60) as response, path.open("wb") as fh:
        fh.write(response.read())


def load_workbook_public() -> object:
    payload = fetch_json(API)
    matches = [f for f in payload.get("files", []) if f.get("name") == TARGET]
    assert len(matches) == 1, matches
    item = matches[0]
    assert int(item["id"]) == 39338195
    assert item.get("computed_md5") == "049db31f8ebcec42c4c44ebe4a6af70a"
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / TARGET
        download(item["download_url"], path)
        assert path.stat().st_size == int(item["size"])
        return load_workbook(path, read_only=False, data_only=True)


def extract_alleles(wb: object):
    ws = wb["Genetic__data"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    assert header[:4] == ["Habitat", "Pop", "Develomental_stage", "ID"], header[:4]
    locus_cols = list(range(4, len(header), 2))
    assert len(locus_cols) == 8

    alleles: dict[tuple[str, str, str, int], list[int]] = defaultdict(list)
    individuals: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    habitat_by_pop: dict[str, str] = {}

    for row in rows[1:]:
        habitat = str(row[0]).strip() if row[0] is not None else ""
        pop = str(row[1]).strip() if row[1] is not None else ""
        stage = str(row[2]).strip() if row[2] is not None else ""
        individual = str(row[3]).strip() if row[3] is not None else ""
        if habitat not in VALID_HABITATS or not pop or stage not in STAGES or not individual:
            continue
        if pop in habitat_by_pop:
            assert habitat_by_pop[pop] == habitat
        habitat_by_pop[pop] = habitat
        individuals[(habitat, pop, stage)].add(individual)
        for locus_i, col in enumerate(locus_cols):
            a1, a2 = number(row[col]), number(row[col + 1])
            if a1 is None or a2 is None:
                continue
            alleles[(habitat, pop, stage, locus_i)].extend([a1, a2])

    assert len(habitat_by_pop) == 6
    assert sum(v == "CON" for v in habitat_by_pop.values()) == 3
    assert sum(v == "FRA" for v in habitat_by_pop.values()) == 3
    return alleles, individuals, habitat_by_pop


def he_from_alleles(values: list[int], unbiased: bool) -> float:
    assert len(values) >= 4 and len(values) % 2 == 0
    counts: dict[int, int] = defaultdict(int)
    for a in values:
        counts[a] += 1
    total = len(values)
    he = 1.0 - sum((count / total) ** 2 for count in counts.values())
    if unbiased:
        n_diploid = total / 2
        he *= (2 * n_diploid) / (2 * n_diploid - 1)
    return he


def mean_locus_he(groups: list[list[int]], unbiased: bool) -> float:
    vals = [he_from_alleles(g, unbiased) for g in groups if g]
    assert len(vals) == 8
    return stats.mean(vals)


def site_values(alleles, habitat_by_pop, unbiased: bool) -> dict[tuple[str, str], float]:
    out = {}
    for pop, habitat in habitat_by_pop.items():
        for stage in STAGES:
            groups = [alleles[(habitat, pop, stage, locus)] for locus in range(8)]
            out[(pop, stage)] = mean_locus_he(groups, unbiased)
    return out


def habitat_pooled_values(alleles, habitat_by_pop, unbiased: bool) -> dict[tuple[str, str], float]:
    out = {}
    for habitat in ("CON", "FRA"):
        pops = [p for p, h in habitat_by_pop.items() if h == habitat]
        for stage in STAGES:
            groups = []
            for locus in range(8):
                merged = []
                for pop in pops:
                    merged.extend(alleles[(habitat, pop, stage, locus)])
                groups.append(merged)
            out[(habitat, stage)] = mean_locus_he(groups, unbiased)
    return out


def habitat_mean_site_values(site: dict[tuple[str, str], float], habitat_by_pop) -> dict[tuple[str, str], float]:
    out = {}
    for habitat in ("CON", "FRA"):
        pops = [p for p, h in habitat_by_pop.items() if h == habitat]
        for stage in STAGES:
            out[(habitat, stage)] = stats.mean(site[(p, stage)] for p in pops)
    return out


def matches(values: dict[tuple[str, str], float]) -> bool:
    return all(round(values[key] + 1e-12, 2) == target for key, target in TABLE2_HE.items())


def hedges_g(fragmented: list[float], reference: list[float]) -> tuple[float, float]:
    n1, n2 = len(fragmented), len(reference)
    assert n1 == n2 == 3
    df = n1 + n2 - 2
    sp = math.sqrt(((n1 - 1) * stats.stdev(fragmented) ** 2 + (n2 - 1) * stats.stdev(reference) ** 2) / df)
    d = (stats.mean(fragmented) - stats.mean(reference)) / sp
    j = 1 - 3 / (4 * df - 1)
    g = j * d
    var = 1 / n1 + 1 / n2 + g * g / (2 * (n1 + n2))
    return g, var


def pearson(x: list[float], y: list[float]) -> float:
    mx, my = stats.mean(x), stats.mean(y)
    dx, dy = [v - mx for v in x], [v - my for v in y]
    return sum(a * b for a, b in zip(dx, dy)) / math.sqrt(sum(a*a for a in dx) * sum(b*b for b in dy))


def centered(values: dict[str, float], habitat_by_pop: dict[str, str], pops: list[str]) -> list[float]:
    means = {
        h: stats.mean(values[p] for p in pops if habitat_by_pop[p] == h)
        for h in ("CON", "FRA")
    }
    return [values[p] - means[habitat_by_pop[p]] for p in pops]


def main() -> None:
    assert RULE.is_file()
    text = RULE.read_text(encoding="utf-8")
    for token in ("reproducibility-only fallback", "H_O", "H_E", "all four", "No estimator may be selected"):
        assert token in text, token

    wb = load_workbook_public()
    alleles, individuals, habitat_by_pop = extract_alleles(wb)
    pops = sorted(habitat_by_pop)

    site_he = site_values(alleles, habitat_by_pop, unbiased=False)
    site_uhe = site_values(alleles, habitat_by_pop, unbiased=True)
    candidates = {
        "pooled_HE": habitat_pooled_values(alleles, habitat_by_pop, unbiased=False),
        "pooled_uHE": habitat_pooled_values(alleles, habitat_by_pop, unbiased=True),
        "mean_site_HE": habitat_mean_site_values(site_he, habitat_by_pop),
        "mean_site_uHE": habitat_mean_site_values(site_uhe, habitat_by_pop),
    }

    matching = []
    for name, values in candidates.items():
        ok = matches(values)
        print(
            "BROSIMUM_HE_RECON "
            f"estimator={name} match={str(ok).lower()} "
            + " ".join(
                f"{h}_{s}={values[(h,s)]:.12f}/target={TABLE2_HE[(h,s)]:.2f}"
                for h, s in (("CON","AD"),("FRA","AD"),("CON","PR"),("FRA","PR"))
            )
        )
        if ok:
            matching.append(name)

    print(f"BROSIMUM_HE_MATCHING estimators={matching!r}")
    assert matching == [], (
        "Brosimum reconciliation state changed: a predeclared H_E estimator now "
        "reproduces all four Table-2 cells. Re-audit prospectively before any effect calculation."
    )

    assert STATUS.is_file(), STATUS
    status = STATUS.read_text(encoding="utf-8")
    for token in (
        "blocked_publication_raw_genotype_reconciliation",
        "0/4 predeclared estimators reproduce all four publication cells",
        "current G_adult-G_offspring coverage: **5/5**",
        "programme increment from Brosimum: **0**",
        "frozen Phase-1 ML002 C/F effects",
    ):
        assert token in status, token

    print(
        "BROSIMUM_GPAIR_RECONCILIATION_BLOCKED "
        "he_estimators_matching=0 coverage_increment=0 audit_pair_coverage=4/5 current_pair_coverage=5/5"
    )


if __name__ == "__main__":
    main()
