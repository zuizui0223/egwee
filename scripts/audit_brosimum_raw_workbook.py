from __future__ import annotations

import json
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ARTICLE_ID = 22130177
TARGET = "Datos-Brosimum.xlsx"
API = f"https://api.figshare.com/v2/articles/{ARTICLE_ID}"


def fetch_json(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": "egwee-public-data-audit/1.0"})
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.load(response)


def download(url: str, path: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "egwee-public-data-audit/1.0"})
    with urllib.request.urlopen(request, timeout=60) as response, path.open("wb") as fh:
        fh.write(response.read())


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\n", " ").replace("\r", " ")
    return text[:120]


def main() -> None:
    payload = fetch_json(API)
    files = payload.get("files", [])
    matches = [f for f in files if f.get("name") == TARGET]
    assert len(matches) == 1, matches
    item = matches[0]
    download_url = item.get("download_url")
    assert download_url, item

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / TARGET
        download(download_url, path)
        assert path.stat().st_size == int(item["size"])

        wb = load_workbook(path, read_only=True, data_only=True)
        print(f"BROSIMUM_WORKBOOK sheets={wb.sheetnames!r}")
        for ws in wb.worksheets:
            print(
                f"BROSIMUM_SHEET name={ws.title!r} rows={ws.max_row} cols={ws.max_column}"
            )
            for row_idx, values in enumerate(ws.iter_rows(values_only=True), start=1):
                if row_idx > 8:
                    break
                rendered = " | ".join(clean(v) for v in values)
                print(f"BROSIMUM_PREVIEW sheet={ws.title!r} row={row_idx}: {rendered}")


if __name__ == "__main__":
    main()
